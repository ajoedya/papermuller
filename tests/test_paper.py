"""Full paper draft, final AI review gate and export. Run: python -m unittest discover -s tests -v"""
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import llm
import paper
import stages

CONFIG = {'DEMO_MODE': True, 'MAX_PAPERS': 12, 'MIN_VERIFIED': 5, 'REVISION_ROUNDS': 3,
          'TARGET_SCORE': 8, 'LANGUAGE': 'English', 'CANDIDATES': 30}
MATRIX = [{'id': 'P1', 'authors': 'A', 'year': 2020, 'title': 'T1', 'doi': '10.0/1', 'main_finding': 'f'},
          {'id': 'P2', 'authors': 'B', 'year': 2021, 'title': 'T2', 'doi': '10.0/2', 'main_finding': 'g'}]
PLAN = {'title': 'T', 'research_questions': ['Q?'], 'scope_note': 'S', 'original_topic': 'topic',
        'author_methods': '', 'author_results': ''}


def sections(**changes):
    base = {'title': 'A title', 'keywords': ['a', 'b'],
            'abstract': ' '.join(['word'] * 160) + ' Results: [pending author analysis]',
            'introduction': 'Intro [P1].', 'discussion': 'Discussion [P2].', 'conclusion': 'Conclusion.'}
    base.update(changes)
    return base


class PaperTests(unittest.TestCase):
    def test_demo_run_catches_invented_result_then_passes(self):
        with tempfile.TemporaryDirectory() as name:
            folder = Path(name)
            plan = stages.stage_plan('road infrastructure governance', CONFIG, folder)
            result = stages.run_research(plan, CONFIG, folder, say=lambda _x: None)
            first = json.loads((folder / '9_paper_review_v1.json').read_text(encoding='utf-8'))
            self.assertFalse(first['passed'])
            self.assertFalse(first['fabrication_gate_passed'])
            final = json.loads((folder / '9_paper_review.json').read_text(encoding='utf-8'))
            self.assertTrue(final['passed'])
            self.assertEqual(final['from_round'], 2)
            self.assertEqual(result['paper_status'], 'ai_review_passed_awaiting_researcher')
            draft = (folder / '9_paper_draft.md').read_text(encoding='utf-8')
            self.assertIn('PASSED AI REVIEW, AWAITING RESEARCHER VERIFICATION', draft)
            self.assertNotIn('12 percent', draft)
            for heading in ('## Abstract', '## 1. Introduction', '## 2. Literature Review', '## 3. Methods',
                            '## 4. Results', '## 5. Discussion', '## 6. Conclusion', '## References'):
                self.assertIn(heading, draft)
            self.assertTrue((folder / '9_paper_draft.docx').exists())
            self.assertFalse((folder / '9_paper_final.md').exists())

    def test_failed_gate_is_stamped_and_cannot_be_released(self):
        bad = {'scores': {k: 9 for k in paper.PAPER_CRITERIA}, 'fixes': [], 'strengths': []}
        bad['scores']['no_fabrication'] = 4
        with tempfile.TemporaryDirectory() as name:
            folder = Path(name)
            plan = stages.stage_plan('road infrastructure governance', CONFIG, folder)
            with patch('paper.ask_for_json', side_effect=lambda prompt, s, stage: (
                    bad if stage == 'paper_review' else llm.ask_for_json(prompt, s, stage=stage))):
                result = stages.run_research(plan, CONFIG, folder, say=lambda _x: None)
            self.assertEqual(result['paper_status'], 'did_not_pass_ai_review')
            self.assertIn('DID NOT PASS AI REVIEW', (folder / '9_paper_draft.md').read_text(encoding='utf-8'))
            self.assertFalse(paper.release_final(folder, say=lambda _x: None))
            self.assertFalse((folder / '9_paper_final.md').exists())

    def test_release_writes_unstamped_final_files(self):
        with tempfile.TemporaryDirectory() as name:
            folder = Path(name)
            plan = stages.stage_plan('road infrastructure governance', CONFIG, folder)
            stages.run_research(plan, CONFIG, folder, say=lambda _x: None)
            self.assertTrue(paper.release_final(folder, say=lambda _x: None))
            final = (folder / '9_paper_final.md').read_text(encoding='utf-8')
            self.assertNotIn('DRAFT —', final)
            self.assertTrue((folder / '9_paper_final.docx').exists())
            review = json.loads((folder / '9_paper_review.json').read_text(encoding='utf-8'))
            self.assertEqual(review['status'], 'released')

    def test_deterministic_checks(self):
        ok = sections()
        md = paper.assemble_paper(ok, PLAN, 'Review [P1; P2].', MATRIX)
        self.assertTrue(paper.paper_checks(ok, md, MATRIX, PLAN)['passed'])
        cited = sections(discussion='Discussion [P9].')
        md = paper.assemble_paper(cited, PLAN, 'Review [P1].', MATRIX)
        self.assertEqual(paper.paper_checks(cited, md, MATRIX, PLAN)['unknown_ids'], ['P9'])
        claims = sections(abstract=' '.join(['word'] * 160) + ' We find a large effect.')
        self.assertFalse(paper.paper_checks(claims, '', MATRIX, PLAN)['passed'])
        leftover = sections(conclusion='[AUTHOR TO WRITE] later')
        self.assertEqual(paper.paper_checks(leftover, '', MATRIX, PLAN)['placeholders'], ['[AUTHOR TO WRITE]'])
        self.assertIn('conclusion', paper.paper_checks(sections(conclusion=''), '', MATRIX, PLAN)['empty_sections'])

    def test_author_methods_and_results_are_copied_unchanged(self):
        plan = dict(PLAN, author_methods='We used **panel** data.', author_results='Effect was 0.4.')
        md = paper.assemble_paper(sections(), plan, 'Review [P1].', MATRIX)
        self.assertIn('We used **panel** data.', md)
        self.assertIn('Effect was 0.4.', md)
        self.assertNotIn('[AUTHOR TO COMPLETE]', md)
        # With author results the abstract is not forced to say results are pending.
        no_pending = sections(abstract=' '.join(['word'] * 160))
        self.assertTrue(paper.paper_checks(no_pending, md, MATRIX, plan)['passed'])

    def test_model_routing_falls_back(self):
        s = {'MODEL': 'base', 'WRITER_MODEL': 'writer', 'REVIEW_MODEL': 'reviewer'}
        self.assertEqual(llm.model_for_stage('paper_draft', s), 'writer')
        self.assertEqual(llm.model_for_stage('paper_review', s), 'reviewer')
        s.update(PAPER_MODEL='pw', PAPER_REVIEW_MODEL='pr')
        self.assertEqual(llm.model_for_stage('paper_draft', s), 'pw')
        self.assertEqual(llm.model_for_stage('paper_review', s), 'pr')

    def test_docx_export_keeps_structure(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as name:
            path = Path(name) / 'p.docx'
            paper.markdown_to_docx('# Title\n\n## Abstract\n\nSome **bold** text.\n\n- one\n- two\n\n> note', path)
            doc = Document(path)
            texts = [p.text for p in doc.paragraphs]
            self.assertEqual(texts[:3], ['Title', 'Abstract', 'Some bold text.'])
            self.assertEqual(doc.paragraphs[0].style.name, 'Title')
            self.assertEqual(doc.paragraphs[3].style.name, 'List Bullet')
            self.assertTrue(any(r.bold for r in doc.paragraphs[2].runs))


if __name__ == '__main__':
    unittest.main()
